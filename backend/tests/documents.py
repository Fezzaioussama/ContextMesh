"""Build real document bytes for parser tests: PDF by hand, Office formats by library."""

from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pptx import Presentation


def pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF; an empty string yields a page with no text layer."""
    font = 3 + 2 * len(pages)
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", _page_tree(len(pages))]
    for index, text in enumerate(pages):
        objects.extend(_page(3 + 2 * index, text, font))
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    return _document(objects)


def _page_tree(count: int) -> bytes:
    kids = " ".join(f"{3 + 2 * index} 0 R" for index in range(count))
    return f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode()


def _page(number: int, text: str, font: int) -> list[bytes]:
    lines = "".join(f"({line}) Tj T* " for line in text.splitlines())
    stream = f"BT /F1 12 Tf 14 TL 72 720 Td {lines}ET".encode() if text else b""
    page = (
        f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {number + 1} 0 R "
        f"/Resources << /Font << /F1 {font} 0 R >> >> >>"
    ).encode()
    content = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
    return [page, content]


def _document(objects: list[bytes]) -> bytes:
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def docx() -> bytes:
    document = Document()
    document.add_heading("Security handbook", 1)
    document.add_paragraph("Rotate keys every quarter.")
    document.add_heading("Access", 2)
    document.add_paragraph("Admins approve new accounts.")
    table = document.add_table(rows=2, cols=2)
    for row, values in enumerate((("Role", "Owner"), ("Admin", "Platform team"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    return _saved(document)


def pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Launch plan"
    slide.placeholders[1].text = "Beta opens in May."
    slide.notes_slide.notes_text_frame.text = "Mention the waitlist."
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Risks"
    second.placeholders[1].text = "Capacity is limited."
    return _saved(presentation)


def xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Item", "Cost"])
    sheet.append(["Servers", 1200])
    sheet.append([None, None])
    sheet.append(["Licenses", 300])
    return _saved(workbook)


def _saved(document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
